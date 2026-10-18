"""Offline phase-one evidence review; never publishes a PIT universe.

Reads the existing CSI cache with its original parsers. Writes only new files.
No price files, network calls, model/factor code, or formal entrypoints are used.
"""
import argparse
from collections import Counter
from dataclasses import asdict
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import sys
from urllib.parse import unquote

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

from qbot.research.universe.browser_observations import read_browser_notice
from qbot.research.universe.notice_attachments import discover_notice_attachments
from qbot.research.universe.notice_pdf import review_csi500_pdf_pages
from qbot.research.universe.notice_review import review_csi500_notice
from qbot.research.universe.notice_workbook import review_csi500_workbook_values
from qbot.research.universe.official_sources import OfficialNoticeSource
from qbot.research.data.schema import normalize_symbol


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, value):
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def identities(changes):
    return {(r["action"], r["symbol"], re.sub(r"\s+", "", r["name"])) for r in changes}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--private-output", required=True)
    args = ap.parse_args()
    repo, output, private = [Path(s).resolve() for s in (args.repo, args.output, args.private_output)]
    if repo not in output.parents or repo not in private.parents:
        raise ValueError("new evidence outputs must remain inside this repository")
    output.mkdir(parents=True, exist_ok=True)
    private.mkdir(parents=True, exist_ok=True)
    source_files, sources, matrix, reviews = {}, {}, [], []

    def register(path, role, url=None, expected=None, representation=None):
        path = Path(path).resolve()
        if not path.is_file() or path.is_symlink():
            raise ValueError("regular evidence file required")
        raw = path.read_bytes()
        digest = sha(raw)
        if expected and digest != expected:
            raise ValueError("existing source hash mismatch: " + str(path))
        rel = str(path.relative_to(repo)).replace("\\", "/")
        source_files[rel] = digest
        key = role + ":" + digest
        sources[key] = {"source_id": key, "role": role, "path": rel, "sha256": digest,
                        "bytes": len(raw), "source_url": url or "UNKNOWN",
                        "representation": representation or "Existing local source bytes; see original receipt",
                        "historical_availability": "UNKNOWN", "financial_admission": False}
        return raw, key

    events_path = repo / "docs/research/evidence/five-year-data-feasibility-v1/pit-events.json"
    events = read_json(events_path)["events"]
    register(events_path, "prior_event_index")
    if len(events) != 14 or len({e["notice_id"] for e in events}) != 14:
        raise ValueError("expected the existing 14 explicitly scoped candidates")
    for event in events:
        number = event["notice_id"]
        ev = event["source_evidence"]
        body_path = repo / ev["local_relative_path"]
        raw, body_id = register(body_path, "notice_body", ev["source_url"],
                                ev.get("raw_content_sha256", ev.get("container_sha256")), ev["representation"])
        if "container_sha256" in ev:
            observation = read_browser_notice(body_path)
            if observation.response_text_sha256 != ev["response_text_utf8_sha256"]:
                raise ValueError("browser response text hash mismatch")
            notice = observation.notice
            body_original_hash = "UNKNOWN"
            body_text_hash = observation.response_text_sha256
        else:
            archive = OfficialNoticeSource(body_path.parent, cache_only=True, max_retries=0)
            notice, artifact = archive.notice(number)
            if archive.network_requests != 0 or artifact.content != raw:
                raise ValueError("offline original-body binding failed")
            body_original_hash, body_text_hash = sha(raw), "NOT_APPLICABLE"
        body = review_csi500_notice(notice)
        if body.publish_date != event["announcement_publish_date"]:
            raise ValueError("previous summary and current source publication disagree")
        attachments, attachment_refs = [], []
        for att in event["attachments"]:
            content, aid = register(repo / att["local_relative_path"], "notice_attachment", att["publisher_url"], att["raw_content_sha256"], "Publisher attachment response bytes retained in prior cache")
            register(repo / att["receipt_relative_path"], "attachment_acquisition_receipt", expected=att["receipt_sha256"])
            receipt = read_json(repo / att["receipt_relative_path"])
            url = receipt["request"]["url"]
            if unquote(url) != unquote(att["publisher_url"]):
                raise ValueError("receipt URL and candidate attachment URL differ")
            cache = OfficialNoticeSource((repo / att["local_relative_path"]).parent, cache_only=True, max_retries=0)
            replay = cache.attachment(url)
            if cache.network_requests != 0 or replay.content != content:
                raise ValueError("attachment cache replay failed")
            attachments.append((att, content, aid))
            attachment_refs.append(aid)
        if len(attachments) != 1:
            raise ValueError("unexpected attachment version scope")
        att, content, aid = attachments[0]
        changes, excluded_reserves, details = [], 0, {}
        if content.startswith(b"%PDF-"):
            reader = PdfReader(io.BytesIO(content), strict=True)
            pages = [p.extract_text() or "" for p in reader.pages]
            parsed = review_csi500_pdf_pages(notice, pages)
            changes = list(parsed.review.changes)
            details = {"parser": "existing notice_pdf.review_csi500_pdf_pages", "pdf_pages": len(pages),
                       "table_heading_page": parsed.heading_page, "pairs_by_page": parsed.pair_counts_by_page,
                       "visual_review_by_this_script": False}
            write_json(private / (str(number) + "-pdf-extraction.json"), {"attachment_sha256": sha(content), "pages": pages, "parsed": asdict(parsed)})
        else:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheets = [{"sheet": sheet.title, "values": [list(row) for row in sheet.iter_rows(values_only=True)]} for sheet in book]
            book.close()
            if {s["sheet"] for s in sheets}.issubset({"调入", "调出", "备选名单"}) and {"调入", "调出"}.issubset({s["sheet"] for s in sheets}):
                # Bounded independent re-read of the already documented split layout.
                for sheet in sheets:
                    values = sheet["values"]
                    if sheet["sheet"] == "备选名单":
                        excluded_reserves += sum(row[:2] == ["000905", "中证500"] for row in values)
                        continue
                    if values[0][:4] != ["指数代码", "指数简称", "证券代码", "证券简称"]:
                        raise ValueError("unrecognized split-workbook header")
                    action = {"调入": "ADD", "调出": "REMOVE"}[sheet["sheet"]]
                    for row_number, row in enumerate(values[1:], 2):
                        if row[:2] != ["000905", "中证500"]:
                            continue
                        code, name = row[2:4]
                        if not isinstance(code, str) or not re.fullmatch(r"\d{6}", code) or not isinstance(name, str):
                            raise ValueError("literal source security code/name required; no padding")
                        changes.append({"action": action, "symbol": normalize_symbol(code), "raw_code": code, "name": name,
                                        "source_sheet": sheet["sheet"], "source_row": row_number,
                                        "source_range": "C%d:D%d" % (row_number, row_number)})
                if body.changes and identities(changes) != identities(body.changes):
                    raise ValueError("native workbook and existing HTML parser disagree")
                details = {"parser": "existing notice_review plus read-only native workbook comparison of documented split layout", "source_sheets": [s["sheet"] for s in sheets], "reserves_excluded": excluded_reserves}
            elif number == 13420:
                # Preserve the explicit '-' instead of guessing China's Energy ticker.
                values = sheets[0]["values"]
                if values[:2] != [["指数代码", "指数简称", "调出", None, "调入", None], [None, None, "证券代码", "证券简称", "证券代码", "证券简称"]]:
                    raise ValueError("conditional workbook header changed")
                for row_number, row in enumerate(values[2:], 3):
                    if row[:2] != ["000905", "中证500"]:
                        continue
                    for action, column in (("REMOVE", 2), ("ADD", 4)):
                        code, name = row[column:column+2]
                        if code != "-" and (not isinstance(code, str) or not re.fullmatch(r"\d{6}", code)):
                            raise ValueError("unrecognized conditional code")
                        changes.append({"action": action, "raw_code": code, "symbol": normalize_symbol(code) if code != "-" else None,
                                        "symbol_status": "SOURCE_LITERAL" if code != "-" else "UNKNOWN_SOURCE_PLACEHOLDER",
                                        "name": name, "source_sheet": sheets[0]["sheet"], "source_row": row_number,
                                        "source_range": "%s%d:%s%d" % (chr(65+column), row_number, chr(66+column), row_number)})
                details = {"parser": "documented six-column conditional source layout; '-' explicitly retained", "production_workbook_parser_acceptance": "NOT_CLAIMED_UNRESOLVED_CODE"}
            else:
                for sheet in sheets:
                    while sheet["values"] and not any(v is not None for v in sheet["values"][-1]):
                        sheet["values"].pop()
                parsed = review_csi500_workbook_values(notice, sheets)
                changes = list(parsed.review.changes)
                details = {"parser": "existing notice_workbook.review_csi500_workbook_values", "source_sheets": [s["sheet"] for s in sheets]}
        counts = Counter(c["action"] for c in changes)
        if counts != {"ADD": event["candidate_add_count"], "REMOVE": event["candidate_remove_count"]}:
            raise ValueError("source candidate count differs from prior evidence: " + str(number))
        if event.get("actual_effective_member_date") is not None:
            raise ValueError("unexpected previously certified actual effective date")
        row = {"notice_id": number, "event_kind": event["event_kind"], "official_url": event["official_notice_url"],
               "announcement_publish_date": body.publish_date, "announcement_publish_time": "UNKNOWN",
               "planned_after_close_date": body.explicit_after_close_date or "UNKNOWN",
               "planned_explicit_start_date": body.explicit_start_date or "UNKNOWN",
               "condition": body.conditional_effective_clause or "NOT_APPLICABLE_IN_THIS_NOTICE",
               "actual_effective_member_date": "UNKNOWN", "actual_effective_confidence": "INSUFFICIENT",
               "historical_known_at": "UNKNOWN", "historical_attachment_vintage": "UNKNOWN",
               "retrieved_at_utc": ev["retrieved_at_utc"],
               "body_original_response_sha256": body_original_hash, "body_response_text_sha256": body_text_hash,
               "body_local_file_sha256": sha(raw), "attachment_original_sha256": sha(content),
               "attachment_url": att["publisher_url"], "attachment_create_at": att.get("publisher_create_at") or "UNKNOWN",
               "attachment_last_modified": att.get("http_last_modified") or "UNKNOWN",
               "candidate_add": counts["ADD"], "candidate_remove": counts["REMOVE"],
               "unresolved_codes": sum(c.get("symbol") is None for c in changes),
               "statement_confidence": "HIGH_FOR_CURRENT_DOCUMENT_STATEMENT_ONLY", "qualification": "CANDIDATE_NOT_PIT",
               "scope_complete": False, "formal_admission": False}
        matrix.append(row)
        reviews.append({"matrix": row, "body_source_id": body_id, "attachment_source_ids": attachment_refs,
                        "parser_details": details, "changes": changes,
                        "pairing_note": "ADD/REMOVE are independent sets; display rows do not prove pairwise economic substitution",
                        "missing": ["independent implementation confirmation", "boundary member snapshots", "historical byte availability", "complete corrections/temporary-event chain"]})

    assert len(matrix) == 14 and sum(r["event_kind"] == "REGULAR" for r in matrix) == 10
    assert sum(r["candidate_add"] for r in matrix if r["event_kind"] == "REGULAR") == 500
    frame = pd.DataFrame(matrix)
    with (output / "pit-evidence-matrix.csv").open("x", encoding="utf-8", newline="") as stream:
        frame.to_csv(stream, index=False)
    write_json(output / "candidate-events.json", {"schema": "PIT_PHASE1_CANDIDATE_EVIDENCE_V1", "no_live_money": True,
               "formal_admission": False, "chain_completeness": "UNKNOWN", "starting_membership_baseline": "UNKNOWN",
               "daily_members_generated": False, "events": reviews})
    write_json(output / "source-index.json", {"schema": "PIT_PHASE1_EXISTING_SOURCE_INDEX_V1", "sources": list(sources.values()), "historical_origin_authentication": "UNKNOWN"})
    write_json(output / "offline-check-receipt.json", {"checked_at_utc": datetime.now(timezone.utc).isoformat(),
               "script_sha256": sha(Path(__file__).read_bytes()), "python": sys.version, "network_requests": 0,
               "market_price_files_opened": 0, "candidate_events": len(matrix), "regular_events": 10,
               "original_file_sha256_at_read": source_files, "actual_dates_remaining_unknown": 14,
               "unknown_source_code_preserved": True, "member_universe_published": False,
               "formal_admission": False, "no_live_money": True})
    print(json.dumps({"events": len(matrix), "regular": 10, "candidate_adds": sum(r["candidate_add"] for r in matrix),
                      "candidate_removes": sum(r["candidate_remove"] for r in matrix), "actual_dates_certified": 0,
                      "bound_files": len(source_files), "output": str(output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
